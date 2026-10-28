"""Build the "my predictions" spreadsheet for the coming season.

Usage: python predictions/make_sheet.py        -> predictions/<season>.xlsx (never overwrites)

Fill the yellow "My wins" column in Excel; the East/West sheets rank themselves (formulas
only, no macros). The Streamlit page "My Predictions" reads the same file. Ties are broken
alphabetically: real tiebreakers (head-to-head…) can't be known from win totals alone.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from db.connection import query_df

YELLOW, GREEN, AMBER, GREY = "FFF2CC", "D9EAD3", "FCE5CD", "EFEFEF"
N = 30  # teams; rows 2..31 on the picks sheet
TOTAL_WINS = 82 * N // 2  # every game has exactly one winner


def build(path: Path) -> None:
    last = query_df("""SELECT full_name, team_abr, conference, division, w, l, season FROM team_seasons
                       WHERE season = (SELECT MAX(season) FROM team_seasons)
                       ORDER BY conference, full_name""")
    assert len(last) == N, f"expected {N} teams, got {len(last)}"
    season = int(last["season"].iat[0])

    wb = Workbook()
    picks = wb.active
    picks.title = "My picks"
    headers = ["Team", "Abbr", "Conference", "Division", f"{season - 1}-{str(season)[2:]} W-L",
               "My wins", "My losses", "Conf. rank", "key"]
    picks.append(headers)
    for r, t in enumerate(last.itertuples(), start=2):
        picks.append([t.full_name, t.team_abr, "East" if t.conference == "E" else "West", t.division,
                      f"{t.w}-{t.l}", None,
                      f'=IF(F{r}="","",82-F{r})',
                      # rank within conference: teams with more wins, then same wins but earlier name
                      f'=IF(F{r}="","",COUNTIFS($C$2:$C$31,C{r},$F$2:$F$31,">"&F{r})'
                      f'+COUNTIFS($C$2:$C$31,C{r},$F$2:$F$31,F{r},$A$2:$A$31,"<"&A{r})+1)',
                      f'=IF(H{r}="","",C{r}&"-"&H{r})'])
        picks[f"F{r}"].fill = PatternFill("solid", fgColor=YELLOW)
    picks.column_dimensions["I"].hidden = True  # lookup key for the standings sheets
    wins_only = DataValidation(type="whole", operator="between", formula1="0", formula2="82",
                               error="Wins must be a whole number from 0 to 82.", showErrorMessage=True)
    picks.add_data_validation(wins_only)
    wins_only.add("F2:F31")

    picks["A33"], picks["F33"] = "Total wins", "=SUM(F2:F31)"
    picks["A34"], picks["F34"] = "Must equal", TOTAL_WINS
    picks["A35"] = "Check"
    picks["F35"] = (f'=IF(COUNT(F2:F31)<{N},{N}-COUNT(F2:F31)&" teams left",'
                    f'IF(F33={TOTAL_WINS},"✓ adds up",IF(F33<{TOTAL_WINS},({TOTAL_WINS}-F33)&" wins short",'
                    f'(F33-{TOTAL_WINS})&" wins too many")))')
    for cell in ("A33", "A34", "A35", "F35"):
        picks[cell].font = Font(bold=True)

    for conference in ("East", "West"):
        sheet = wb.create_sheet(conference)
        sheet.append(["Seed", "Team", "W", "L", "GB", "Status"])
        for seed in range(1, 16):
            r = seed + 1
            match = f'MATCH("{conference}-"&A{r},\'My picks\'!$I$2:$I$31,0)'
            sheet.append([seed,
                          f"=IFERROR(INDEX('My picks'!$A$2:$A$31,{match}),\"\")",
                          f"=IFERROR(INDEX('My picks'!$F$2:$F$31,{match}),\"\")",
                          f'=IF(C{r}="","",82-C{r})',
                          f'=IF(C{r}="","",$C$2-C{r})',
                          f'=IF(B{r}="","",IF(A{r}<=6,"Playoffs",IF(A{r}<=10,"Play-in","Lottery")))'])
        for status, color in (("Playoffs", GREEN), ("Play-in", AMBER), ("Lottery", GREY)):
            sheet.conditional_formatting.add(
                "A2:F16", FormulaRule(formula=[f'$F2="{status}"'], fill=PatternFill("solid", fgColor=color)))
        sheet.column_dimensions["B"].width = 26

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    picks.column_dimensions["A"].width = 26
    picks.column_dimensions["D"].width = 12
    picks.column_dimensions["E"].width = 12
    wb.save(path)


if __name__ == "__main__":
    season = int(query_df("SELECT MAX(season) + 1 FROM team_seasons").iat[0, 0])
    out = Path(__file__).parent / f"{season - 1}-{str(season)[2:]}.xlsx"  # e.g. 2026-27.xlsx
    if out.exists():
        raise SystemExit(f"{out} already exists: delete it first if you really want a blank sheet")
    build(out)
    print(f"Created {out}: fill the yellow 'My wins' column in Excel")
